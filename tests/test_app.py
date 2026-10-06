import sys
import os
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

import pytest
from datetime import datetime, date, timedelta

from app import create_app
from app.extensions import db
from app.models import User, Client, Location, ClientContact, Asset, MaintenanceLog, MicrosoftLicense, Ticket, TicketIntervention, SystemSetting, MaintenanceEvent, MaintenanceEventItem

from app.services.license_notifier import check_and_notify_license_expirations

@pytest.fixture
def app():
    app = create_app('test')
    with app.app_context():
        db.create_all()
        yield app
        db.session.remove()
        db.drop_all()

@pytest.fixture
def client(app):
    return app.test_client()

def test_user_authentication(app):
    with app.app_context():
        user = User(name='Test Tech', email='tech@test.com', role='tech')
        user.set_password('Secret123!')
        db.session.add(user)
        db.session.commit()

        fetched_user = User.query.filter_by(email='tech@test.com').first()
        assert fetched_user is not None
        assert fetched_user.check_password('Secret123!') is True
        assert fetched_user.check_password('WrongPass') is False
        assert fetched_user.is_admin is False

def test_client_and_location_creation(app):
    with app.app_context():
        cli = Client(
            company_name='Empresa Prueba SAS',
            tax_id='999.888.777-1',
            contact_name='Juan Pérez',
            contact_email='juan@prueba.com',
            contact_phone='3001112233',
            address='Calle 123'
        )
        db.session.add(cli)
        db.session.commit()

        loc = Location(client_id=cli.id, name='Sede Norte', city='Bogotá')
        db.session.add(loc)
        db.session.commit()

        assert len(cli.locations) == 1
        assert cli.locations[0].name == 'Sede Norte'

def test_asset_and_maintenance(app):
    with app.app_context():
        user = User(name='Tech 1', email='t1@test.com', role='tech')
        user.set_password('pass')
        cli = Client(company_name='Cli Asset', tax_id='12345', contact_name='A', contact_email='a@a.com', contact_phone='123')
        db.session.add_all([user, cli])
        db.session.commit()

        asset = Asset(
            internal_code='EQ-TEST-001',
            device_type='Servidor',
            brand='Dell',
            model='PowerEdge',
            status='Operativo',
            client_id=cli.id
        )
        db.session.add(asset)
        db.session.commit()

        maint = MaintenanceLog(
            asset_id=asset.id,
            user_id=user.id,
            log_type='Preventivo',
            description='Limpieza de polvo y revisión de RAM',
            cost=50.0
        )
        db.session.add(maint)
        db.session.commit()

        assert len(asset.maintenance_logs) == 1
        assert asset.maintenance_logs[0].log_type == 'Preventivo'

def test_license_expiration_alerts(app):
    with app.app_context():
        admin = User(name='Admin', email='admin@test.com', role='admin')
        admin.set_password('admin')
        cli = Client(company_name='Cli Lic', tax_id='777', contact_name='B', contact_email='b@b.com', contact_phone='123')
        db.session.add_all([admin, cli])
        db.session.commit()

        today = date.today()
        lic_critical = MicrosoftLicense(
            client_id=cli.id,
            subscription_type='M365 Business Premium',
            quantity=10,
            acquisition_date=today - timedelta(days=360),
            expiration_date=today + timedelta(days=5) # 5 days left => critical
        )
        lic_active = MicrosoftLicense(
            client_id=cli.id,
            subscription_type='Azure Plan',
            quantity=1,
            acquisition_date=today - timedelta(days=100),
            expiration_date=today + timedelta(days=90) # 90 days left => active
        )
        db.session.add_all([lic_critical, lic_active])
        db.session.commit()

        assert lic_critical.alert_status == 'critical'
        assert lic_active.alert_status == 'active'

        alerts = check_and_notify_license_expirations()
        assert len(alerts) >= 1
        assert alerts[0]['subscription'] == 'M365 Business Premium'

def test_license_rf07_fields(app):
    with app.app_context():
        cli = Client(company_name='Cli RF07', tax_id='999000', contact_name='X', contact_email='x@x.com', contact_phone='123')
        db.session.add(cli)
        db.session.commit()

        today = date.today()
        lic = MicrosoftLicense(
            client_id=cli.id,
            subscription_type='Office 365 E3',
            quantity=5,
            acquisition_date=today,
            expiration_date=today + timedelta(days=365),
            installation_user='admin@tenant.com',
            installation_password='Password123!',
            activation_code='AAAAA-BBBBB-CCCCC-DDDDD-EEEEE',
            download_link='https://setup.office.com',
            assigned_user_names='Juan Pérez, María Gómez'
        )
        db.session.add(lic)
        db.session.commit()

        fetched = MicrosoftLicense.query.filter_by(subscription_type='Office 365 E3').first()
        assert fetched.installation_user == 'admin@tenant.com'
        assert fetched.installation_password == 'Password123!'
        assert fetched.activation_code == 'AAAAA-BBBBB-CCCCC-DDDDD-EEEEE'
        assert fetched.download_link == 'https://setup.office.com'
        assert fetched.assigned_user_names == 'Juan Pérez, María Gómez'

def test_license_excel_export(client, app):
    with app.app_context():
        user = User(name='Admin Export', email='admin_export@test.com', role='admin')
        user.set_password('pass')
        cli = Client(company_name='Empresa Export SAS', tax_id='123456789', contact_name='Cont', contact_email='c@exp.com', contact_phone='123')
        db.session.add_all([user, cli])
        db.session.commit()

        today = date.today()
        lic = MicrosoftLicense(
            client_id=cli.id,
            subscription_type='M365 Business Standard',
            quantity=15,
            acquisition_date=today,
            expiration_date=today + timedelta(days=180),
            installation_user='admin@empresaexport.com',
            installation_password='SecretPassword123',
            activation_code='EXCEL-TEST-CODE-12345',
            download_link='https://setup.office.com',
            assigned_user_names='Usuario 1, Usuario 2',
            notes='Licencia adquirida para departamento TI'
        )
        db.session.add(lic)
        db.session.commit()

        client.post('/login', data={'email': 'admin_export@test.com', 'password': 'pass'}, follow_redirects=True)

        response = client.get('/licenses/export/excel')
        assert response.status_code == 200
        assert response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        assert 'reporte_licencias_' in response.headers.get('Content-Disposition', '')



def test_ticket_sla_calculation(app):
    with app.app_context():
        cli = Client(company_name='Cli Ticket', tax_id='888', contact_name='C', contact_email='c@c.com', contact_phone='123')
        db.session.add(cli)
        db.session.commit()

        start_time = datetime.utcnow()
        tck_critical = Ticket(
            ticket_code='TCK-2026-TEST1',
            client_id=cli.id,
            title='Caída de Servidor Principal',
            description='Urgente',
            priority='Crítica',
            created_at=start_time,
            sla_due_at=Ticket.calculate_sla_due('Crítica', start_time)
        )
        db.session.add(tck_critical)
        db.session.commit()

        expected_due = start_time + timedelta(hours=2)
        assert abs((tck_critical.sla_due_at - expected_due).total_seconds()) < 5
        assert tck_critical.is_sla_breached is False

def test_dashboard_rendering(client, app):
    with app.app_context():
        user = User(name='Admin User', email='admin@test.com', role='admin')
        user.set_password('pass')
        db.session.add(user)
        db.session.commit()

        # Login
        client.post('/login', data={'email': 'admin@test.com', 'password': 'pass'}, follow_redirects=True)
        response = client.get('/')
        assert response.status_code == 200
        assert b'Panel de Control Operativo' in response.data

def test_client_contacts_and_ticket_notifications(app):
    with app.app_context():
        cli = Client(company_name='Cli Contact Test', tax_id='55443322', contact_name='P', contact_email='p@p.com', contact_phone='123')
        db.session.add(cli)
        db.session.commit()

        contact = ClientContact(
            client_id=cli.id,
            name='Carlos Restrepo',
            email='crestrepo@client.com',
            phone='3001234567',
            position='Gerente TI',
            is_primary=True
        )
        db.session.add(contact)
        db.session.commit()

        assert len(cli.contacts) == 1
        assert cli.contacts[0].email == 'crestrepo@client.com'

        # Ticket with contact
        tck = Ticket(
            ticket_code='TCK-2026-CONT1',
            client_id=cli.id,
            contact_id=contact.id,
            title='Solicitud de Acceso',
            description='Acceso a carpeta compartida',
            priority='Baja',
            created_at=datetime.utcnow(),
            sla_due_at=Ticket.calculate_sla_due('Baja')
        )
        db.session.add(tck)
        db.session.commit()

        assert tck.contact is not None
        assert tck.contact.name == 'Carlos Restrepo'

def test_system_settings(app):
    with app.app_context():
        from app.models import SystemSetting
        settings = SystemSetting.get_settings()
        assert settings.mail_server == 'smtp.gmail.com'
        assert settings.notifications_enabled is True

        settings.mail_server = 'smtp.office365.com'
        db.session.commit()

        fetched = SystemSetting.get_settings()
        assert fetched.mail_server == 'smtp.office365.com'

def test_edit_user(client, app):
    with app.app_context():
        admin = User(name='Admin User', email='admin_edit@test.com', role='admin')
        admin.set_password('pass')
        tech = User(name='Tech Edit', email='tech_edit@test.com', role='tech')
        tech.set_password('pass')
        db.session.add_all([admin, tech])
        db.session.commit()
        tech_id = tech.id

        client.post('/login', data={'email': 'admin_edit@test.com', 'password': 'pass'}, follow_redirects=True)

        response = client.post(f'/technicians/{tech_id}/edit', data={
            'name': 'Tech Editado',
            'email': 'tech_edited@test.com',
            'role': 'admin',
            'password': 'NewPass123!'
        }, follow_redirects=True)
        assert response.status_code == 200

        updated = db.session.get(User, tech_id)
        assert updated.name == 'Tech Editado'
        assert updated.email == 'tech_edited@test.com'
        assert updated.role == 'admin'
        assert updated.check_password('NewPass123!') is True

        # Password empty should keep the existing one
        client.post(f'/technicians/{tech_id}/edit', data={
            'name': 'Tech Editado',
            'email': 'tech_edited@test.com',
            'role': 'admin',
            'password': ''
        }, follow_redirects=True)
        unchanged = db.session.get(User, tech_id)
        assert unchanged.check_password('NewPass123!') is True

def test_edit_user_duplicate_email(client, app):
    with app.app_context():
        admin = User(name='Admin User', email='admin_dup@test.com', role='admin')
        admin.set_password('pass')
        tech1 = User(name='Tech One', email='tech_one@test.com', role='tech')
        tech1.set_password('pass')
        tech2 = User(name='Tech Two', email='tech_two@test.com', role='tech')
        tech2.set_password('pass')
        db.session.add_all([admin, tech1, tech2])
        db.session.commit()
        tech1_id = tech1.id

        client.post('/login', data={'email': 'admin_dup@test.com', 'password': 'pass'}, follow_redirects=True)

        response = client.post(f'/technicians/{tech1_id}/edit', data={
            'name': 'Tech One',
            'email': 'tech_two@test.com',
            'role': 'tech',
            'password': ''
        }, follow_redirects=True)
        assert response.status_code == 200

        unchanged = db.session.get(User, tech1_id)
        assert unchanged.email == 'tech_one@test.com'

def test_toggle_user_status(client, app):
    with app.app_context():
        admin = User(name='Admin User', email='admin_toggle@test.com', role='admin')
        admin.set_password('pass')
        tech = User(name='Tech Toggle', email='tech_toggle@test.com', role='tech')
        tech.set_password('pass')
        db.session.add_all([admin, tech])
        db.session.commit()
        tech_id = tech.id

        client.post('/login', data={'email': 'admin_toggle@test.com', 'password': 'pass'}, follow_redirects=True)

        # Deactivate
        response = client.post(f'/technicians/{tech_id}/toggle-status', follow_redirects=True)
        assert response.status_code == 200
        assert db.session.get(User, tech_id).is_active is False

        # Reactivate
        client.post(f'/technicians/{tech_id}/toggle-status', follow_redirects=True)
        assert db.session.get(User, tech_id).is_active is True

def test_toggle_own_status_blocked(client, app):
    with app.app_context():
        admin = User(name='Admin Self', email='admin_self@test.com', role='admin')
        admin.set_password('pass')
        db.session.add(admin)
        db.session.commit()
        admin_id = admin.id

        client.post('/login', data={'email': 'admin_self@test.com', 'password': 'pass'}, follow_redirects=True)

        client.post(f'/technicians/{admin_id}/toggle-status', follow_redirects=True)
        assert db.session.get(User, admin_id).is_active is True

def test_edit_location(client, app):
    with app.app_context():
        user = User(name='Admin Loc', email='admin_loc@test.com', role='admin')
        user.set_password('pass')
        cli = Client(company_name='Cli Loc', tax_id='CL-001', contact_name='X', contact_email='x@x.com', contact_phone='123')
        db.session.add_all([user, cli])
        db.session.commit()

        loc = Location(client_id=cli.id, name='Sede Norte', city='Bogotá', address='Calle 1')
        db.session.add(loc)
        db.session.commit()
        cli_id = cli.id
        loc_id = loc.id

        client.post('/login', data={'email': 'admin_loc@test.com', 'password': 'pass'}, follow_redirects=True)

        response = client.post(f'/clients/{cli_id}/edit-location/{loc_id}', data={
            'name': 'Sede Sur',
            'city': 'Medellín',
            'address': 'Carrera 80',
            'contact_person': 'Ana López',
            'contact_phone': '3009998877'
        }, follow_redirects=True)
        assert response.status_code == 200

        updated = db.session.get(Location, loc_id)
        assert updated.name == 'Sede Sur'
        assert updated.city == 'Medellín'
        assert updated.address == 'Carrera 80'
        assert updated.contact_person == 'Ana López'
        assert updated.contact_phone == '3009998877'

def test_edit_location_empty_name(client, app):
    with app.app_context():
        user = User(name='Admin Loc2', email='admin_loc2@test.com', role='admin')
        user.set_password('pass')
        cli = Client(company_name='Cli Loc2', tax_id='CL-002', contact_name='Y', contact_email='y@y.com', contact_phone='123')
        db.session.add_all([user, cli])
        db.session.commit()
        loc = Location(client_id=cli.id, name='Sede Original')
        db.session.add(loc)
        db.session.commit()
        cli_id, loc_id = cli.id, loc.id

        client.post('/login', data={'email': 'admin_loc2@test.com', 'password': 'pass'}, follow_redirects=True)

        client.post(f'/clients/{cli_id}/edit-location/{loc_id}', data={'name': ''}, follow_redirects=True)
        unchanged = db.session.get(Location, loc_id)
        assert unchanged.name == 'Sede Original'

def test_delete_location_unlinks_assets_and_tickets(client, app):
    with app.app_context():
        user = User(name='Admin Loc3', email='admin_loc3@test.com', role='admin')
        user.set_password('pass')
        cli = Client(company_name='Cli Loc3', tax_id='CL-003', contact_name='Z', contact_email='z@z.com', contact_phone='123')
        db.session.add_all([user, cli])
        db.session.commit()

        loc = Location(client_id=cli.id, name='Sede a Borrar')
        db.session.add(loc)
        db.session.commit()

        asset = Asset(internal_code='EQ-LOC-1', device_type='PC', brand='HP', model='X', status='Operativo', client_id=cli.id, location_id=loc.id)
        tck = Ticket(ticket_code='TCK-LOC-1', client_id=cli.id, location_id=loc.id, title='T', description='D', priority='Baja', created_at=datetime.utcnow(), sla_due_at=Ticket.calculate_sla_due('Baja'))
        db.session.add_all([asset, tck])
        db.session.commit()
        cli_id, loc_id, asset_id, tck_id = cli.id, loc.id, asset.id, tck.id

        client.post('/login', data={'email': 'admin_loc3@test.com', 'password': 'pass'}, follow_redirects=True)

        response = client.post(f'/clients/{cli_id}/delete-location/{loc_id}', follow_redirects=True)
        assert response.status_code == 200

        assert db.session.get(Location, loc_id) is None
        assert db.session.get(Asset, asset_id).location_id is None
        assert db.session.get(Ticket, tck_id).location_id is None

def test_delete_asset(client, app):
    with app.app_context():
        user = User(name='Admin Del Asset', email='admin_del_asset@test.com', role='admin')
        user.set_password('pass')
        cli = Client(company_name='Cli Del Asset', tax_id='DA-001', contact_name='X', contact_email='x@x.com', contact_phone='123')
        db.session.add_all([user, cli])
        db.session.commit()

        asset = Asset(internal_code='EQ-DEL-1', device_type='PC', brand='HP', model='X', status='Operativo', client_id=cli.id)
        db.session.add(asset)
        db.session.commit()

        log = MaintenanceLog(asset_id=asset.id, user_id=user.id, log_type='Preventivo', description='Limpieza', cost=10)
        tck = Ticket(ticket_code='TCK-DEL-1', client_id=cli.id, asset_id=asset.id, title='T', description='D', priority='Baja', created_at=datetime.utcnow(), sla_due_at=Ticket.calculate_sla_due('Baja'))
        db.session.add_all([log, tck])
        db.session.commit()
        asset_id, tck_id = asset.id, tck.id

        client.post('/login', data={'email': 'admin_del_asset@test.com', 'password': 'pass'}, follow_redirects=True)

        response = client.post(f'/assets/{asset_id}/delete', follow_redirects=True)
        assert response.status_code == 200

        assert db.session.get(Asset, asset_id) is None
        assert db.session.get(Ticket, tck_id).asset_id is None

def test_edit_contact(client, app):
    with app.app_context():
        user = User(name='Admin Contact', email='admin_contact@test.com', role='admin')
        user.set_password('pass')
        cli = Client(company_name='Cli Contact Edit', tax_id='CE-001', contact_name='X', contact_email='x@x.com', contact_phone='123')
        db.session.add_all([user, cli])
        db.session.commit()

        contact = ClientContact(client_id=cli.id, name='Carlos', email='carlos@empresa.com', phone='3001234567', position='Soporte', is_primary=False)
        db.session.add(contact)
        db.session.commit()
        cli_id, contact_id = cli.id, contact.id

        client.post('/login', data={'email': 'admin_contact@test.com', 'password': 'pass'}, follow_redirects=True)

        response = client.post(f'/clients/{cli_id}/edit-contact/{contact_id}', data={
            'name': 'Carlos Actualizado',
            'email': 'carlos.nuevo@empresa.com',
            'phone': '3009998877',
            'position': 'Gerente',
            'is_primary': '1'
        }, follow_redirects=True)
        assert response.status_code == 200

        updated = db.session.get(ClientContact, contact_id)
        assert updated.name == 'Carlos Actualizado'
        assert updated.email == 'carlos.nuevo@empresa.com'
        assert updated.phone == '3009998877'
        assert updated.position == 'Gerente'
        assert updated.is_primary is True

def test_edit_contact_sets_only_one_primary(client, app):
    with app.app_context():
        user = User(name='Admin Contact2', email='admin_contact2@test.com', role='admin')
        user.set_password('pass')
        cli = Client(company_name='Cli Contact 2', tax_id='CE-002', contact_name='X', contact_email='x@x.com', contact_phone='123')
        db.session.add_all([user, cli])
        db.session.commit()

        c1 = ClientContact(client_id=cli.id, name='Uno', email='uno@e.com', is_primary=True)
        c2 = ClientContact(client_id=cli.id, name='Dos', email='dos@e.com', is_primary=False)
        db.session.add_all([c1, c2])
        db.session.commit()
        cli_id, c2_id = cli.id, c2.id

        client.post('/login', data={'email': 'admin_contact2@test.com', 'password': 'pass'}, follow_redirects=True)

        client.post(f'/clients/{cli_id}/edit-contact/{c2_id}', data={
            'name': 'Dos',
            'email': 'dos@e.com',
            'phone': '',
            'position': '',
            'is_primary': '1'
        }, follow_redirects=True)

        primaries = ClientContact.query.filter_by(client_id=cli_id, is_primary=True).all()
        assert len(primaries) == 1
        assert primaries[0].id == c2_id

def test_tickets_excel_export(client, app):
    with app.app_context():
        user = User(name='Admin Tickets', email='admin_tickets@test.com', role='admin')
        user.set_password('pass')
        cli1 = Client(company_name='Cliente Export Tickets', tax_id='TEX-001', contact_name='X', contact_email='x@x.com', contact_phone='123')
        cli2 = Client(company_name='Otro Cliente', tax_id='TEX-002', contact_name='Y', contact_email='y@y.com', contact_phone='123')
        db.session.add_all([user, cli1, cli2])
        db.session.commit()

        now = datetime.utcnow()
        t1 = Ticket(ticket_code='TCK-EXP-0001', client_id=cli1.id, title='Ticket A', description='D', priority='Alta', created_at=now, sla_due_at=Ticket.calculate_sla_due('Alta', now))
        t2 = Ticket(ticket_code='TCK-EXP-0002', client_id=cli1.id, title='Ticket B', description='D', priority='Baja', created_at=now, sla_due_at=Ticket.calculate_sla_due('Baja', now))
        t3 = Ticket(ticket_code='TCK-EXP-0003', client_id=cli2.id, title='Ticket C', description='D', priority='Media', created_at=now - timedelta(days=30), sla_due_at=Ticket.calculate_sla_due('Media', now - timedelta(days=30)))
        db.session.add_all([t1, t2, t3])
        db.session.commit()
        cli1_id = cli1.id

        client.post('/login', data={'email': 'admin_tickets@test.com', 'password': 'pass'}, follow_redirects=True)

        # Export sin filtros
        response = client.get('/tickets/export/excel')
        assert response.status_code == 200
        assert response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        assert 'reporte_tickets_' in response.headers.get('Content-Disposition', '')

        # Export filtrado por cliente
        response_cli = client.get(f'/tickets/export/excel?client_id={cli1_id}')
        assert response_cli.status_code == 200

        import openpyxl
        from io import BytesIO
        wb = openpyxl.load_workbook(BytesIO(response_cli.data))
        ws = wb.active
        # Fila 1 = encabezados, luego 2 tickets del cliente 1 (no el del cliente 2)
        assert ws.max_row == 3
        cliente_col = [ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)]
        assert all('Cliente Export Tickets' in (c or '') for c in cliente_col)

def test_tickets_excel_export_date_filter(client, app):
    with app.app_context():
        user = User(name='Admin Tickets2', email='admin_tickets2@test.com', role='admin')
        user.set_password('pass')
        cli = Client(company_name='Cli Fecha', tax_id='TEX-003', contact_name='X', contact_email='x@x.com', contact_phone='123')
        db.session.add(user)
        db.session.add(cli)
        db.session.commit()

        today = datetime.utcnow()
        old = today - timedelta(days=40)
        t_recent = Ticket(ticket_code='TCK-DATE-0001', client_id=cli.id, title='Reciente', description='D', priority='Media', created_at=today, sla_due_at=Ticket.calculate_sla_due('Media', today))
        t_old = Ticket(ticket_code='TCK-DATE-0002', client_id=cli.id, title='Antiguo', description='D', priority='Media', created_at=old, sla_due_at=Ticket.calculate_sla_due('Media', old))
        db.session.add_all([t_recent, t_old])
        db.session.commit()
        cli_id = cli.id
        date_from = (today - timedelta(days=7)).strftime('%Y-%m-%d')

        client.post('/login', data={'email': 'admin_tickets2@test.com', 'password': 'pass'}, follow_redirects=True)

        response = client.get(f'/tickets/export/excel?client_id={cli_id}&date_from={date_from}')
        assert response.status_code == 200

        import openpyxl
        from io import BytesIO
        wb = openpyxl.load_workbook(BytesIO(response.data))
        ws = wb.active
        # Solo el ticket reciente
        assert ws.max_row == 2
        assert ws.cell(row=2, column=2).value == 'TCK-DATE-0001'

def test_assets_excel_export(client, app):
    with app.app_context():
        user = User(name='Admin Assets', email='admin_assets_exp@test.com', role='admin')
        user.set_password('pass')
        cli1 = Client(company_name='Cliente Equipos', tax_id='AEX-001', contact_name='X', contact_email='x@x.com', contact_phone='123')
        cli2 = Client(company_name='Otro Cliente Eq', tax_id='AEX-002', contact_name='Y', contact_email='y@y.com', contact_phone='123')
        db.session.add_all([user, cli1, cli2])
        db.session.commit()

        a1 = Asset(internal_code='EQ-EXP-0001', device_type='PC Desktop', brand='Dell', model='OptiPlex', status='Operativo', client_id=cli1.id)
        a2 = Asset(internal_code='EQ-EXP-0002', device_type='Laptop', brand='HP', model='ProBook', status='En Mantenimiento', client_id=cli1.id)
        a3 = Asset(internal_code='EQ-EXP-0003', device_type='Servidor', brand='Lenovo', model='ThinkSystem', status='Operativo', client_id=cli2.id)
        db.session.add_all([a1, a2, a3])
        db.session.commit()

        log = MaintenanceLog(asset_id=a1.id, user_id=user.id, log_type='Preventivo', description='Limpieza', cost=25)
        db.session.add(log)
        db.session.commit()
        cli1_id = cli1.id

        client.post('/login', data={'email': 'admin_assets_exp@test.com', 'password': 'pass'}, follow_redirects=True)

        # Export sin filtros
        response = client.get('/assets/export/excel')
        assert response.status_code == 200
        assert response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        assert 'reporte_hojas_de_vida_' in response.headers.get('Content-Disposition', '')

        # Export filtrado por cliente
        response_cli = client.get(f'/assets/export/excel?client_id={cli1_id}')
        assert response_cli.status_code == 200

        import openpyxl
        from io import BytesIO
        wb = openpyxl.load_workbook(BytesIO(response_cli.data))
        ws = wb.active
        # Fila 1 = encabezados, luego 2 equipos del cliente 1
        assert ws.max_row == 3
        cliente_col = [ws.cell(row=r, column=4).value for r in range(2, ws.max_row + 1)]
        assert all('Cliente Equipos' in (c or '') for c in cliente_col)

def test_assets_excel_export_date_filter(client, app):
    with app.app_context():
        user = User(name='Admin Assets2', email='admin_assets_exp2@test.com', role='admin')
        user.set_password('pass')
        cli = Client(company_name='Cli Fecha Eq', tax_id='AEX-003', contact_name='X', contact_email='x@x.com', contact_phone='123')
        db.session.add_all([user, cli])
        db.session.commit()

        today = datetime.utcnow()
        old = today - timedelta(days=40)
        a_recent = Asset(internal_code='EQ-DATE-0001', device_type='PC', brand='Dell', model='X', status='Operativo', client_id=cli.id, created_at=today)
        a_old = Asset(internal_code='EQ-DATE-0002', device_type='PC', brand='HP', model='Y', status='Operativo', client_id=cli.id, created_at=old)
        db.session.add_all([a_recent, a_old])
        db.session.commit()
        cli_id = cli.id
        date_from = (today - timedelta(days=7)).strftime('%Y-%m-%d')

        client.post('/login', data={'email': 'admin_assets_exp2@test.com', 'password': 'pass'}, follow_redirects=True)

        response = client.get(f'/assets/export/excel?client_id={cli_id}&date_from={date_from}')
        assert response.status_code == 200

        import openpyxl
        from io import BytesIO
        wb = openpyxl.load_workbook(BytesIO(response.data))
        ws = wb.active
        # Solo el equipo reciente
        assert ws.max_row == 2
        assert ws.cell(row=2, column=2).value == 'EQ-DATE-0001'


def test_maintenance_event_create_with_multiple_assets(client, app):
    with app.app_context():
        user = User(name='Admin Mant', email='admin_mant@test.com', role='admin')
        user.set_password('pass')
        cli = Client(company_name='Cli Mant', tax_id='MAN-001', contact_name='X', contact_email='x@x.com', contact_phone='123')
        db.session.add_all([user, cli])
        db.session.commit()

        a1 = Asset(internal_code='MAN-EQ-1', device_type='PC', brand='HP', model='A', status='Operativo', client_id=cli.id)
        a2 = Asset(internal_code='MAN-EQ-2', device_type='Laptop', brand='Dell', model='B', status='Operativo', client_id=cli.id)
        db.session.add_all([a1, a2])
        db.session.commit()
        cid, a1id, a2id = cli.id, a1.id, a2.id

        client.post('/login', data={'email': 'admin_mant@test.com', 'password': 'pass'}, follow_redirects=True)

        response = client.post('/maintenance/create', data={
            'client_id': cid,
            'maintenance_type': 'Preventivo',
            'observations': 'Prueba de evento',
            'asset_ids': [a1id, a2id],
        }, follow_redirects=True)
        assert response.status_code == 200

        event = MaintenanceEvent.query.first()
        assert event is not None
        assert event.status == 'Abierto'
        assert event.client_id == cid
        assert len(event.items) == 2

def test_maintenance_event_status_finish_sets_date(client, app):
    with app.app_context():
        user = User(name='Admin Mant2', email='admin_mant2@test.com', role='admin')
        user.set_password('pass')
        cli = Client(company_name='Cli Mant2', tax_id='MAN-002', contact_name='X', contact_email='x@x.com', contact_phone='123')
        db.session.add_all([user, cli])
        db.session.commit()

        a1 = Asset(internal_code='MAN-EQ-3', device_type='PC', brand='HP', model='A', status='Operativo', client_id=cli.id)
        db.session.add(a1)
        db.session.commit()

        event = MaintenanceEvent(client_id=cli.id, technician_id=user.id, status='Abierto')
        db.session.add(event)
        db.session.commit()
        db.session.add(MaintenanceEventItem(event_id=event.id, asset_id=a1.id))
        db.session.commit()
        eid = event.id

        client.post('/login', data={'email': 'admin_mant2@test.com', 'password': 'pass'}, follow_redirects=True)

        # Marcar como terminado -> debe registrar finished_at
        resp = client.post(f'/maintenance/{eid}/update-status',
                           data={'status': 'Terminado', 'observations': 'Finalizado'},
                           follow_redirects=True)
        assert resp.status_code == 200

        updated = db.session.get(MaintenanceEvent, eid)
        assert updated.status == 'Terminado'
        assert updated.finished_at is not None

        # Reabrir -> debe limpiar finished_at
        client.post(f'/maintenance/{eid}/update-status',
                    data={'status': 'En elaboración'},
                    follow_redirects=True)
        reopened = db.session.get(MaintenanceEvent, eid)
        assert reopened.status == 'En elaboración'
        assert reopened.finished_at is None

def test_maintenance_upload_photos(client, app, tmp_path):
    with app.app_context():
        from io import BytesIO
        user = User(name='Admin Mant3', email='admin_mant3@test.com', role='admin')
        user.set_password('pass')
        cli = Client(company_name='Cli Mant3', tax_id='MAN-003', contact_name='X', contact_email='x@x.com', contact_phone='123')
        db.session.add_all([user, cli])
        db.session.commit()

        a1 = Asset(internal_code='MAN-EQ-4', device_type='PC', brand='HP', model='A', status='Operativo', client_id=cli.id)
        db.session.add(a1)
        db.session.commit()

        event = MaintenanceEvent(client_id=cli.id, technician_id=user.id, status='Abierto')
        db.session.add(event)
        db.session.commit()
        item = MaintenanceEventItem(event_id=event.id, asset_id=a1.id)
        db.session.add(item)
        db.session.commit()
        eid, iid = event.id, item.id

        client.post('/login', data={'email': 'admin_mant3@test.com', 'password': 'pass'}, follow_redirects=True)

        data = {
            'photo_before': (BytesIO(b'fake-before'), 'antes.jpg'),
            'photo_after': (BytesIO(b'fake-after'), 'despues.jpg'),
            'notes': 'Todo ok',
        }
        resp = client.post(f'/maintenance/{eid}/item/{iid}/photos',
                           data=data, content_type='multipart/form-data', follow_redirects=True)
        assert resp.status_code == 200

        updated = db.session.get(MaintenanceEventItem, iid)
        assert updated.photo_before is not None
        assert updated.photo_after is not None
        assert updated.notes == 'Todo ok'

        # Limpieza de archivos generados durante la prueba
        import os as _os
        from flask import current_app
        upload_dir = current_app.config['UPLOAD_FOLDER']
        for fn in (updated.photo_before, updated.photo_after):
            path = _os.path.join(upload_dir, fn)
            if _os.path.exists(path):
                _os.remove(path)


def test_asset_bitacora_fed_by_finished_events_only(client, app):
    """La bitácora del equipo (req-03, 2.3) solo muestra eventos TERMINADOS."""
    with app.app_context():
        user = User(name='Admin Bit', email='admin_bit@test.com', role='admin')
        user.set_password('pass')
        cli = Client(company_name='Cli Bit', tax_id='BIT-001', contact_name='X', contact_email='x@x.com', contact_phone='123')
        db.session.add_all([user, cli])
        db.session.commit()

        asset = Asset(internal_code='BIT-EQ-1', device_type='PC', brand='HP', model='A', status='Operativo', client_id=cli.id)
        db.session.add(asset)
        db.session.commit()

        # Evento ABIERTO -> no debe aparecer en la bitácora
        ev_open = MaintenanceEvent(client_id=cli.id, technician_id=user.id, status='Abierto')
        db.session.add(ev_open)
        db.session.commit()
        db.session.add(MaintenanceEventItem(event_id=ev_open.id, asset_id=asset.id))
        db.session.commit()

        # Evento TERMINADO -> sí debe aparecer
        ev_done = MaintenanceEvent(client_id=cli.id, technician_id=user.id, status='Terminado',
                                   maintenance_type='Preventivo', observations='Listo',
                                   finished_at=datetime.utcnow())
        db.session.add(ev_done)
        db.session.commit()
        db.session.add(MaintenanceEventItem(event_id=ev_done.id, asset_id=asset.id, notes='Equipo ok'))
        db.session.commit()

        db.session.refresh(asset)
        records = asset.maintenance_records
        assert len(records) == 1
        assert records[0].event.id == ev_done.id
        assert records[0].event.status == 'Terminado'

        # La vista del detalle renderiza la bitácora y NO tiene el botón de registrar mantenimiento
        client.post('/login', data={'email': 'admin_bit@test.com', 'password': 'pass'}, follow_redirects=True)
        resp = client.get(f'/assets/{asset.id}')
        assert resp.status_code == 200
        html = resp.data.decode('utf-8')
        assert 'Bitácora de Mantenimientos' in html
        # La ruta de mantenimiento del asset ya no debe existir en la página
        assert f'/assets/{asset.id}/add-maintenance' not in html
        # El evento terminado debe reflejarse en la bitácora
        assert f'Evento #{ev_done.id}' in html


def test_maintenance_event_pdf_report(client, app):
    """El informe PDF del evento (req-03, 2.4) se genera correctamente."""
    with app.app_context():
        user = User(name='Admin PDF', email='admin_pdf@test.com', role='admin')
        user.set_password('pass')
        cli = Client(company_name='Cliente PDF SAS', tax_id='PDF-001', contact_name='X', contact_email='x@x.com', contact_phone='123')
        db.session.add_all([user, cli])
        db.session.commit()

        loc = Location(client_id=cli.id, name='Sede PDF', city='Bogotá')
        db.session.add(loc)
        db.session.commit()

        a1 = Asset(internal_code='EQ-PDF-1', device_type='Servidor', brand='Dell', model='PowerEdge',
                   cpu='Xeon', ram='32GB', storage='1TB', os_installed='Windows Server',
                   status='Operativo', client_id=cli.id, location_id=loc.id, serial_number='SN-PDF-1')
        a2 = Asset(internal_code='EQ-PDF-2', device_type='Laptop', brand='HP', model='ProBook',
                   status='Operativo', client_id=cli.id)
        db.session.add_all([a1, a2])
        db.session.commit()

        event = MaintenanceEvent(client_id=cli.id, technician_id=user.id, status='Terminado',
                                 maintenance_type='Preventivo', observations='Mantenimiento general.',
                                 finished_at=datetime.utcnow())
        db.session.add(event)
        db.session.commit()
        db.session.add_all([
            MaintenanceEventItem(event_id=event.id, asset_id=a1.id, notes='Limpieza de ventiladores.'),
            MaintenanceEventItem(event_id=event.id, asset_id=a2.id, notes='Actualización de drivers.'),
        ])
        db.session.commit()
        eid = event.id

        client.post('/login', data={'email': 'admin_pdf@test.com', 'password': 'pass'}, follow_redirects=True)

        response = client.get(f'/maintenance/{eid}/export/pdf')
        assert response.status_code == 200
        assert response.mimetype == 'application/pdf'
        assert response.data[:4] == b'%PDF'
        assert 'informe_mantenimiento_evento_' in response.headers.get('Content-Disposition', '')


def test_maintenance_event_pdf_with_photos(client, app):
    """El informe PDF incluye las fotos antes/después cuando existen."""
    with app.app_context():
        from PIL import Image
        import os
        upload_dir = app.config['UPLOAD_FOLDER']
        os.makedirs(upload_dir, exist_ok=True)
        before_name = 'test_pdf_before.png'
        after_name = 'test_pdf_after.png'
        Image.new('RGB', (400, 300), (30, 80, 160)).save(os.path.join(upload_dir, before_name))
        Image.new('RGB', (400, 300), (20, 150, 90)).save(os.path.join(upload_dir, after_name))

        user = User(name='Admin PDF2', email='admin_pdf2@test.com', role='admin')
        user.set_password('pass')
        cli = Client(company_name='Cliente Fotos PDF', tax_id='PDF-002', contact_name='X', contact_email='x@x.com', contact_phone='123')
        db.session.add_all([user, cli])
        db.session.commit()

        a1 = Asset(internal_code='EQ-PDF-3', device_type='PC', brand='HP', model='X',
                   status='Operativo', client_id=cli.id, serial_number='SN-PDF-3')
        db.session.add(a1)
        db.session.commit()

        event = MaintenanceEvent(client_id=cli.id, technician_id=user.id, status='Terminado',
                                 maintenance_type='Correctivo', finished_at=datetime.utcnow())
        db.session.add(event)
        db.session.commit()
        db.session.add(MaintenanceEventItem(event_id=event.id, asset_id=a1.id, notes='Reparado',
                                            photo_before=before_name, photo_after=after_name))
        db.session.commit()
        eid = event.id

        client.post('/login', data={'email': 'admin_pdf2@test.com', 'password': 'pass'}, follow_redirects=True)

        try:
            response = client.get(f'/maintenance/{eid}/export/pdf')
            assert response.status_code == 200
            assert response.data[:4] == b'%PDF'
            # Un PDF con imágenes embebidas es más grande que uno sin ellas
            assert len(response.data) > 4500
        finally:
            for fn in (before_name, after_name):
                path = os.path.join(upload_dir, fn)
                if os.path.exists(path):
                    os.remove(path)


def test_asset_add_maintenance_route_removed(client, app):
    """La ruta de registrar mantenimiento desde hojas de vida fue eliminada (req-03, 2.3)."""
    with app.app_context():
        user = User(name='Admin NoMaint', email='admin_nomaint@test.com', role='admin')
        user.set_password('pass')
        cli = Client(company_name='Cli NoMaint', tax_id='NM-001', contact_name='X', contact_email='x@x.com', contact_phone='123')
        db.session.add_all([user, cli])
        db.session.commit()

        asset = Asset(internal_code='NM-EQ-1', device_type='PC', brand='HP', model='A', status='Operativo', client_id=cli.id)
        db.session.add(asset)
        db.session.commit()
        aid = asset.id

        client.post('/login', data={'email': 'admin_nomaint@test.com', 'password': 'pass'}, follow_redirects=True)

        resp = client.post(f'/assets/{aid}/add-maintenance', data={
            'log_type': 'Preventivo', 'description': 'X'
        })
        # Debe responder 404 (ruta inexistente)
        assert resp.status_code == 404
