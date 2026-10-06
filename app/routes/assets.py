from datetime import datetime
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from flask import Blueprint, render_template, redirect, url_for, flash, request, jsonify, send_file
from flask_login import login_required
from app.extensions import db
from app.models.asset import Asset
from app.models.client import Client

assets_bp = Blueprint('assets', __name__)

@assets_bp.route('/')
@login_required
def index():
    query = request.args.get('q', '').strip()
    client_id = request.args.get('client_id', type=int)
    device_type = request.args.get('device_type', '').strip()
    status = request.args.get('status', '').strip()
    date_from_str = request.args.get('date_from', '').strip()
    date_to_str = request.args.get('date_to', '').strip()

    assets_query = Asset.query
    if query:
        assets_query = assets_query.filter(
            (Asset.internal_code.ilike(f'%{query}%')) |
            (Asset.serial_number.ilike(f'%{query}%')) |
            (Asset.brand.ilike(f'%{query}%')) |
            (Asset.model.ilike(f'%{query}%'))
        )
    if client_id:
        assets_query = assets_query.filter_by(client_id=client_id)
    if device_type:
        assets_query = assets_query.filter_by(device_type=device_type)
    if status:
        assets_query = assets_query.filter_by(status=status)

    # Filtros por rango de fechas sobre la fecha de registro (created_at)
    if date_from_str:
        try:
            date_from = datetime.strptime(date_from_str, '%Y-%m-%d')
            assets_query = assets_query.filter(Asset.created_at >= date_from)
        except (ValueError, TypeError):
            pass
    if date_to_str:
        try:
            date_to = datetime.strptime(date_to_str, '%Y-%m-%d')
            date_to_end = date_to.replace(hour=23, minute=59, second=59, microsecond=999999)
            assets_query = assets_query.filter(Asset.created_at <= date_to_end)
        except (ValueError, TypeError):
            pass

    assets_list = assets_query.order_by(Asset.created_at.desc()).all()
    clients = Client.query.filter_by(is_active=True).order_by(Client.company_name.asc()).all()
    
    device_types = ['Servidor', 'PC Desktop', 'Laptop', 'Switch', 'Router', 'Firewall', 'Impresora', 'UPS', 'Otro']

    return render_template('assets/index.html',
                           assets=assets_list,
                           clients=clients,
                           device_types=device_types,
                           query=query,
                           client_id=client_id,
                           device_type=device_type,
                           status=status,
                           date_from=date_from_str,
                           date_to=date_to_str)

@assets_bp.route('/create', methods=['GET', 'POST'])
@login_required
def create():
    clients = Client.query.filter_by(is_active=True).order_by(Client.company_name.asc()).all()
    if request.method == 'POST':
        internal_code = request.form.get('internal_code', '').strip()
        serial_number = request.form.get('serial_number', '').strip()
        device_type = request.form.get('device_type', '').strip()
        brand = request.form.get('brand', '').strip()
        model = request.form.get('model', '').strip()
        cpu = request.form.get('cpu', '').strip()
        ram = request.form.get('ram', '').strip()
        storage = request.form.get('storage', '').strip()
        os_installed = request.form.get('os_installed', '').strip()
        status = request.form.get('status', 'Operativo').strip()
        client_id = request.form.get('client_id', type=int)
        location_id = request.form.get('location_id', type=int)
        notes = request.form.get('notes', '').strip()

        if Asset.query.filter_by(internal_code=internal_code).first():
            flash(f'El código interno "{internal_code}" ya existe.', 'danger')
            return render_template('assets/form.html', asset=None, clients=clients)

        new_asset = Asset(
            internal_code=internal_code,
            serial_number=serial_number,
            device_type=device_type,
            brand=brand,
            model=model,
            cpu=cpu,
            ram=ram,
            storage=storage,
            os_installed=os_installed,
            status=status,
            client_id=client_id,
            location_id=location_id if location_id else None,
            notes=notes
        )
        db.session.add(new_asset)
        db.session.commit()
        flash(f'Equipo "{internal_code}" registrado exitosamente en Hoja de Vida.', 'success')
        return redirect(url_for('assets.detail', asset_id=new_asset.id))

    return render_template('assets/form.html', asset=None, clients=clients)

@assets_bp.route('/<int:asset_id>')
@login_required
def detail(asset_id):
    asset = Asset.query.get_or_404(asset_id)
    return render_template('assets/detail.html', asset=asset)

@assets_bp.route('/<int:asset_id>/edit', methods=['GET', 'POST'])
@login_required
def edit(asset_id):
    asset = Asset.query.get_or_404(asset_id)
    clients = Client.query.filter_by(is_active=True).order_by(Client.company_name.asc()).all()
    
    if request.method == 'POST':
        asset.internal_code = request.form.get('internal_code', '').strip()
        asset.serial_number = request.form.get('serial_number', '').strip()
        asset.device_type = request.form.get('device_type', '').strip()
        asset.brand = request.form.get('brand', '').strip()
        asset.model = request.form.get('model', '').strip()
        asset.cpu = request.form.get('cpu', '').strip()
        asset.ram = request.form.get('ram', '').strip()
        asset.storage = request.form.get('storage', '').strip()
        asset.os_installed = request.form.get('os_installed', '').strip()
        asset.status = request.form.get('status', 'Operativo').strip()
        asset.client_id = request.form.get('client_id', type=int)
        
        loc_id = request.form.get('location_id', type=int)
        asset.location_id = loc_id if loc_id else None
        asset.notes = request.form.get('notes', '').strip()

        db.session.commit()
        flash(f'Equipo "{asset.internal_code}" actualizado.', 'success')
        return redirect(url_for('assets.detail', asset_id=asset.id))

    return render_template('assets/form.html', asset=asset, clients=clients)

@assets_bp.route('/<int:asset_id>/delete', methods=['POST'])
@login_required
def delete_asset(asset_id):
    asset = Asset.query.get_or_404(asset_id)
    internal_code = asset.internal_code

    # Desvincular tickets asociados para evitar referencias rotas
    # (los registros de mantenimiento se eliminan en cascada)
    for ticket in asset.tickets:
        ticket.asset_id = None

    db.session.delete(asset)
    db.session.commit()
    flash(f'Equipo "{internal_code}" eliminado de la Hoja de Vida.', 'info')
    return redirect(url_for('assets.index'))

@assets_bp.route('/api/<int:client_id>/assets')
@login_required
def api_assets(client_id):
    assets = Asset.query.filter_by(client_id=client_id).all()
    return jsonify([{'id': a.id, 'internal_code': a.internal_code, 'name': f"{a.internal_code} - {a.device_type} ({a.brand} {a.model})"} for a in assets])


@assets_bp.route('/export/excel')
@login_required
def export_excel():
    client_id = request.args.get('client_id', type=int)
    device_type = request.args.get('device_type', '').strip()
    status = request.args.get('status', '').strip()
    date_from_str = request.args.get('date_from', '').strip()
    date_to_str = request.args.get('date_to', '').strip()

    assets_query = Asset.query
    if client_id:
        assets_query = assets_query.filter_by(client_id=client_id)
    if device_type:
        assets_query = assets_query.filter_by(device_type=device_type)
    if status:
        assets_query = assets_query.filter_by(status=status)

    # Filtros por rango de fechas sobre la fecha de registro (created_at)
    if date_from_str:
        try:
            date_from = datetime.strptime(date_from_str, '%Y-%m-%d')
            assets_query = assets_query.filter(Asset.created_at >= date_from)
        except (ValueError, TypeError):
            pass
    if date_to_str:
        try:
            date_to = datetime.strptime(date_to_str, '%Y-%m-%d')
            date_to_end = date_to.replace(hour=23, minute=59, second=59, microsecond=999999)
            assets_query = assets_query.filter(Asset.created_at <= date_to_end)
        except (ValueError, TypeError):
            pass

    assets_list = assets_query.order_by(Asset.created_at.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hojas de Vida de Equipos"

    headers = [
        "ID",
        "Código Interno",
        "Número de Serial",
        "Cliente",
        "NIT / ID Fiscal",
        "Sede / Ubicación",
        "Tipo de Equipo",
        "Marca",
        "Modelo",
        "Procesador (CPU)",
        "Memoria RAM",
        "Disco / Storage",
        "Sistema Operativo",
        "Estado",
        "N° Mantenimientos (Bitácora)",
        "Tipo Último Mantenimiento",
        "Fecha Último Mantenimiento",
        "Técnico Último Mantenimiento",
        "Notas",
        "Fecha de Registro"
    ]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws.append(headers)
    ws.row_dimensions[1].height = 26

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for a in assets_list:
        # Bitácora alimentada por los eventos de mantenimiento terminados (req-03, 2.3)
        records = a.maintenance_records
        last_record = records[0] if records else None
        last_event = last_record.event if last_record else None
        last_date = None
        if last_event:
            last_date = last_event.finished_at or last_event.opened_at
        row_data = [
            a.id,
            a.internal_code or '',
            a.serial_number or '',
            a.client.company_name if a.client else '',
            a.client.tax_id if a.client else '',
            a.location.name if a.location else 'Sede Principal',
            a.device_type or '',
            a.brand or '',
            a.model or '',
            a.cpu or '',
            a.ram or '',
            a.storage or '',
            a.os_installed or '',
            a.status or '',
            len(records),
            last_event.maintenance_type if last_event else '',
            last_date.strftime('%Y-%m-%d %H:%M') if last_date else '',
            last_event.technician.name if (last_event and last_event.technician) else '',
            a.notes or '',
            a.created_at.strftime('%Y-%m-%d %H:%M') if a.created_at else ''
        ]
        ws.append(row_data)

    for row in range(2, len(assets_list) + 2):
        ws.row_dimensions[row].height = 20
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"reporte_hojas_de_vida_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )
