from datetime import datetime
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from flask import Blueprint, render_template, redirect, url_for, flash, request, send_file
from flask_login import login_required
from app.extensions import db
from app.models.license import MicrosoftLicense
from app.models.client import Client
from app.services.license_notifier import check_and_notify_license_expirations

licenses_bp = Blueprint('licenses', __name__)

@licenses_bp.route('/')
@login_required
def index():
    client_id = request.args.get('client_id', type=int)
    urgency_filter = request.args.get('urgency', '').strip()

    licenses_query = MicrosoftLicense.query
    if client_id:
        licenses_query = licenses_query.filter_by(client_id=client_id)

    licenses_list = licenses_query.all()

    # Filter in python for property-based urgency if specified
    if urgency_filter:
        licenses_list = [l for l in licenses_list if l.alert_status == urgency_filter]

    # Sort by days until expiration ascending
    licenses_list.sort(key=lambda x: x.days_until_expiration)

    clients = Client.query.filter_by(is_active=True).order_by(Client.company_name.asc()).all()

    # Summary counts
    all_licenses = MicrosoftLicense.query.all()
    status_summary = {
        'expired': sum(1 for l in all_licenses if l.alert_status == 'expired'),
        'critical': sum(1 for l in all_licenses if l.alert_status == 'critical'),
        'warning': sum(1 for l in all_licenses if l.alert_status == 'warning'),
        'notice': sum(1 for l in all_licenses if l.alert_status == 'notice'),
        'active': sum(1 for l in all_licenses if l.alert_status == 'active'),
    }

    return render_template('licenses/index.html',
                           licenses=licenses_list,
                           clients=clients,
                           client_id=client_id,
                           urgency_filter=urgency_filter,
                           status_summary=status_summary)

@licenses_bp.route('/create', methods=['GET', 'POST'])
@login_required
def create():
    clients = Client.query.filter_by(is_active=True).order_by(Client.company_name.asc()).all()

    if request.method == 'POST':
        client_id = request.form.get('client_id', type=int)
        subscription_type = request.form.get('subscription_type', '').strip()
        quantity = request.form.get('quantity', 1, type=int)
        assigned_users_count = request.form.get('assigned_users_count', 0, type=int)
        acq_date_str = request.form.get('acquisition_date')
        exp_date_str = request.form.get('expiration_date')
        installation_user = request.form.get('installation_user', '').strip()
        installation_password = request.form.get('installation_password', '').strip()
        activation_code = request.form.get('activation_code', '').strip()
        download_link = request.form.get('download_link', '').strip()
        assigned_user_names = request.form.get('assigned_user_names', '').strip()
        notes = request.form.get('notes', '').strip()

        try:
            acq_date = datetime.strptime(acq_date_str, '%Y-%m-%d').date()
            exp_date = datetime.strptime(exp_date_str, '%Y-%m-%d').date()
        except (ValueError, TypeError):
            flash('Formato de fecha inválido. Utilice AAAA-MM-DD.', 'danger')
            return render_template('licenses/form.html', license=None, clients=clients)

        new_license = MicrosoftLicense(
            client_id=client_id,
            subscription_type=subscription_type,
            quantity=quantity,
            assigned_users_count=assigned_users_count,
            acquisition_date=acq_date,
            expiration_date=exp_date,
            installation_user=installation_user,
            installation_password=installation_password,
            activation_code=activation_code,
            download_link=download_link,
            assigned_user_names=assigned_user_names,
            notes=notes
        )
        db.session.add(new_license)
        db.session.commit()

        flash(f'Licencia "{subscription_type}" registrada exitosamente.', 'success')
        return redirect(url_for('licenses.index'))

    return render_template('licenses/form.html', license=None, clients=clients)

@licenses_bp.route('/<int:license_id>/edit', methods=['GET', 'POST'])
@login_required
def edit(license_id):
    lic = MicrosoftLicense.query.get_or_404(license_id)
    clients = Client.query.filter_by(is_active=True).order_by(Client.company_name.asc()).all()

    if request.method == 'POST':
        lic.client_id = request.form.get('client_id', type=int)
        lic.subscription_type = request.form.get('subscription_type', '').strip()
        lic.quantity = request.form.get('quantity', 1, type=int)
        lic.assigned_users_count = request.form.get('assigned_users_count', 0, type=int)
        acq_date_str = request.form.get('acquisition_date')
        exp_date_str = request.form.get('expiration_date')
        lic.installation_user = request.form.get('installation_user', '').strip()
        lic.installation_password = request.form.get('installation_password', '').strip()
        lic.activation_code = request.form.get('activation_code', '').strip()
        lic.download_link = request.form.get('download_link', '').strip()
        lic.assigned_user_names = request.form.get('assigned_user_names', '').strip()
        lic.notes = request.form.get('notes', '').strip()

        try:
            lic.acquisition_date = datetime.strptime(acq_date_str, '%Y-%m-%d').date()
            lic.expiration_date = datetime.strptime(exp_date_str, '%Y-%m-%d').date()
        except (ValueError, TypeError):
            flash('Formato de fecha inválido.', 'danger')
            return render_template('licenses/form.html', license=lic, clients=clients)

        db.session.commit()
        flash('Licencia actualizada correctamente.', 'success')
        return redirect(url_for('licenses.index'))


    return render_template('licenses/form.html', license=lic, clients=clients)

@licenses_bp.route('/trigger-alerts', methods=['POST'])
@login_required
def trigger_alerts():
    alerts = check_and_notify_license_expirations()
    if alerts:
        flash(f'Proceso de verificación ejecutado: {len(alerts)} alertas detectadas y notificadas.', 'warning')
    else:
        flash('Proceso de verificación ejecutado: No se encontraron licencias críticas con vencimiento próximo.', 'success')
    return redirect(url_for('licenses.index'))


@licenses_bp.route('/export/excel')
@login_required
def export_excel():
    client_id = request.args.get('client_id', type=int)
    urgency_filter = request.args.get('urgency', '').strip()

    licenses_query = MicrosoftLicense.query
    if client_id:
        licenses_query = licenses_query.filter_by(client_id=client_id)

    licenses_list = licenses_query.all()

    if urgency_filter:
        licenses_list = [l for l in licenses_list if l.alert_status == urgency_filter]

    licenses_list.sort(key=lambda x: x.days_until_expiration)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Licencias Microsoft"

    headers = [
        "ID",
        "Cliente",
        "NIT / ID Fiscal",
        "Suscripción Microsoft",
        "Cantidad Contratada",
        "Cantidad Asignada",
        "Fecha de Adquisición",
        "Fecha de Vencimiento",
        "Días Restantes",
        "Estado Alerta",
        "Usuario de Instalación",
        "Clave / Contraseña",
        "Código de Activación",
        "Link de Descarga",
        "Usuarios Asignados",
        "Notas",
        "Fecha Registro"
    ]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws.append(headers)
    ws.row_dimensions[1].height = 26

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for lic in licenses_list:
        row_data = [
            lic.id,
            lic.client.company_name if lic.client else '',
            lic.client.tax_id if lic.client else '',
            lic.subscription_type or '',
            lic.quantity or 0,
            lic.assigned_users_count or 0,
            lic.acquisition_date.strftime('%Y-%m-%d') if lic.acquisition_date else '',
            lic.expiration_date.strftime('%Y-%m-%d') if lic.expiration_date else '',
            lic.days_until_expiration,
            lic.alert_status.upper() if lic.alert_status else '',
            lic.installation_user or '',
            lic.installation_password or '',
            lic.activation_code or '',
            lic.download_link or '',
            lic.assigned_user_names or '',
            lic.notes or '',
            lic.created_at.strftime('%Y-%m-%d %H:%M') if lic.created_at else ''
        ]
        ws.append(row_data)

    for row in range(2, len(licenses_list) + 2):
        ws.row_dimensions[row].height = 20
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"reporte_licencias_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )

